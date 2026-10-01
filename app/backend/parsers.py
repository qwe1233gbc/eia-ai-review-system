"""报告文本解析：TXT / MD / DOCX / PDF → 纯文本。"""
from __future__ import annotations

import io
from pathlib import Path

SUPPORTED_EXT = {".txt", ".md", ".markdown", ".docx", ".pdf"}


def _decode_bytes(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(filename: str, data: bytes) -> str:
    """按扩展名分发解析，返回纯文本。"""
    suffix = Path(filename).suffix.lower()
    if suffix in (".txt", ".md", ".markdown"):
        return _decode_bytes(data)
    if suffix == ".docx":
        return _extract_docx(data)
    if suffix == ".pdf":
        return _extract_pdf(data)
    # 未知类型：尝试当文本解析
    return _decode_bytes(data)


def _extract_docx(data: bytes) -> str:
    import docx  # python-docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages)